# -*- coding: utf-8 -*-
"""Cobertura por estudiante de las evaluaciones AJ Áncash 2026 (TOV + HSE + PRONABEC + COAR).
Salida: Cobertura_Evaluaciones_AJ_Ancash_2026.xlsx + cobertura_estudiantes_ancash.json
"""
import sys, os, json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "pipeline"))
from db import conn

BASE = os.path.dirname(os.path.abspath(__file__))
CONFIG = os.environ.get("AJ_CONFIG", os.path.join(os.path.dirname(BASE), "config", "evaluaciones.json"))
TOVJSON = os.environ.get("AJ_TOV_JSON", os.path.join(os.path.dirname(BASE), "datos",
                                                     "dashboard_tov_resultados.json"))

_EV = json.load(open(CONFIG, encoding="utf-8"))
INSTR = [
    ("TOV (Proceso)",     _EV["tov_proceso"]),
    ("HSE Inicio",        _EV["hse_inicio"]),
    ("HSE Proceso",       _EV["hse_proceso"]),
    ("PRONABEC Comunicación", _EV["pronabec_comunicacion"]),
    ("PRONABEC Matemática",   _EV["pronabec_matematica"]),
    ("COAR",              _EV["coar"]),
]

def main():
    c = conn(); cur = c.cursor()
    # claves (pk del estudiante) de Áncash, desde la vista oficial
    cur.execute("SELECT DISTINCT clave FROM reportes.vw_tov_areas_largo_ancash_2026")
    keys = set(r[0] for r in cur.fetchall())
    # nombres desde el maestro de estudiantes
    cur.execute("""SELECT pk::text, nombres_completo FROM bd_evaluaciones.estudiantes
                   WHERE pk::text = ANY(%s)""", (list(keys),))
    est = {pk: (nom or "(sin nombre)") for pk, nom in cur.fetchall()}
    for k in keys:
        est.setdefault(k, "(sin nombre)")
    print("Estudiantes Áncash (TOV Proceso):", len(keys))
    presentes = {}
    for label, pk in INSTR:
        cur.execute("""SELECT DISTINCT e.fk_evaluado::text FROM bd_evaluaciones.evaluaciones e
                       WHERE e.fk_caracteristica_evaluacion=%s AND e.respuestas <> '{}'::jsonb""", (pk,))
        presentes[label] = set(r[0] for r in cur.fetchall())
        print(f"  {label:22s} {len(presentes[label] & keys)}")
    cur.close(); c.close()

    cols = [l for l, _ in INSTR]
    rows = []
    for k in sorted(keys, key=lambda x: est.get(x, "")):
        rows.append([est[k]] + [("✔" if k in presentes[l] else "—") for l in cols])

    # JSON
    out = {est.get(k, k): {l: (k in presentes[l]) for l in cols} for k in keys}
    json.dump(out, open(os.path.join(BASE, "cobertura_estudiantes_ancash.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)

    # XLSX
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Cobertura"
    hdr = ["Estudiante"] + cols
    ws.append(hdr)
    for r in rows: ws.append(r)
    fill = PatternFill("solid", fgColor="1F4E79")
    for i, h in enumerate(hdr, 1):
        cc = ws.cell(1, i); cc.font = Font(bold=True, color="FFFFFF"); cc.fill = fill
        cc.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 38
    for i in range(2, len(cols) + 2): ws.column_dimensions[chr(64 + i)].width = 18
    ws.freeze_panes = "B2"
    ws2 = wb.create_sheet("Resumen")
    ws2.append(["Evaluación", "Estudiantes con resultado (Áncash)"])
    for l in cols: ws2.append([l, len(presentes[l] & keys)])
    ws2.append(["Total estudiantes", len(keys)])
    for cc in ws2["A1:B1"][0]: cc.font = Font(bold=True, color="FFFFFF"); cc.fill = fill
    out_xlsx = os.path.join(BASE, "Cobertura_Evaluaciones_AJ_Ancash_2026.xlsx")
    wb.save(out_xlsx)
    print("OK:", out_xlsx, "| filas:", len(rows))

if __name__ == "__main__":
    main()
